import io
import logging
import os
from typing import List
from docx import Document
from app.services.files.chunker import DocumentChunker
from app.services.files.models import DocumentChunk, ExtractionResult
from app.services.files.processors.base import BaseFileProcessor

logger = logging.getLogger("solix.files.docx")


class DocxProcessor(BaseFileProcessor):
    """Extracts text, headings, lists, and tables from Word (.docx) documents."""

    def __init__(self):
        self.chunker = DocumentChunker(chunk_size=1000, chunk_overlap=150)

    def can_process(self, category: str, mime_type: str, filename: str) -> bool:
        return category == "docx" or filename.lower().endswith(".docx")

    async def extract(self, file_id: str, filename: str, content_bytes: bytes) -> ExtractionResult:
        stream = io.BytesIO(content_bytes)
        try:
            doc = Document(stream)
        except Exception as exc:
            logger.error(f"[DOCX] Failed to load docx document: {exc}")
            raise ValueError(f"Corrupted or invalid DOCX document: {exc}")

        extracted_elements: List[str] = []
        current_section = "Document Introduction"
        section_texts: List[str] = []
        all_chunks: List[DocumentChunk] = []

        def flush_section():
            nonlocal section_texts, current_section
            if section_texts:
                combined = "\n\n".join(section_texts).strip()
                if combined:
                    sec_chunks = self.chunker.chunk_text(
                        combined,
                        file_id=file_id,
                        filename=filename,
                        section=current_section,
                    )
                    all_chunks.extend(sec_chunks)
            section_texts = []

        # Extract paragraphs with heading detection
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style and p.style.name else ""
            if "heading" in style_name:
                flush_section()
                current_section = text
                extracted_elements.append(f"\n## {text}\n")
            elif "list" in style_name:
                extracted_elements.append(f"- {text}")
                section_texts.append(f"- {text}")
            else:
                extracted_elements.append(text)
                section_texts.append(text)

        # Extract tables
        for table_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                # Avoid repeated merged cells
                cleaned_cells = []
                for c in cells:
                    if not cleaned_cells or c != cleaned_cells[-1]:
                        cleaned_cells.append(c)
                if cleaned_cells:
                    table_rows.append(" | ".join(cleaned_cells))

            if table_rows:
                table_str = f"\n[Table {table_idx + 1}]\n" + "\n".join(table_rows) + "\n"
                extracted_elements.append(table_str)
                section_texts.append(table_str)

        flush_section()

        full_text = "\n\n".join(extracted_elements)

        # Document core properties
        metadata = {}
        try:
            core = doc.core_properties
            if core.title:
                metadata["title"] = core.title
            if core.author:
                metadata["author"] = core.author
            if core.created:
                metadata["created"] = str(core.created)
        except Exception:
            pass

        return ExtractionResult(
            text=full_text,
            chunks=all_chunks,
            metadata=metadata,
        )

