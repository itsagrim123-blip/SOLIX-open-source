import csv
import io
import logging
import os
import xml.etree.ElementTree as ET
import zipfile
from typing import List
import openpyxl
import xlrd
from app.services.files.chunker import DocumentChunker
from app.services.files.models import DocumentChunk, ExtractionResult
from app.services.files.processors.base import BaseFileProcessor

logger = logging.getLogger("solix.files.spreadsheet")


class SpreadsheetProcessor(BaseFileProcessor):
    """Extracts structured tables, columns, and headers from XLSX, XLS, CSV, TSV, and ODS."""

    def __init__(self):
        self.chunker = DocumentChunker(chunk_size=1200, chunk_overlap=100)

    def can_process(self, category: str, mime_type: str, filename: str) -> bool:
        ext = os.path.splitext(filename.lower())[1]
        return category == "spreadsheet" or ext in (".xlsx", ".xls", ".ods", ".csv", ".tsv")

    async def extract(self, file_id: str, filename: str, content_bytes: bytes) -> ExtractionResult:
        ext = os.path.splitext(filename.lower())[1]

        if ext == ".csv" or ext == ".tsv":
            return self._extract_csv(file_id, filename, content_bytes, delimiter="\t" if ext == ".tsv" else ",")
        elif ext == ".xlsx":
            return self._extract_xlsx(file_id, filename, content_bytes)
        elif ext == ".xls":
            return self._extract_xls(file_id, filename, content_bytes)
        elif ext == ".ods":
            return self._extract_ods(file_id, filename, content_bytes)
        else:
            return self._extract_csv(file_id, filename, content_bytes)

    def _extract_csv(self, file_id: str, filename: str, content_bytes: bytes, delimiter: str = ",") -> ExtractionResult:
        # Decode CSV safely
        text = ""
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                text = content_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            text = content_bytes.decode("utf-8", errors="replace")

        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = list(reader)

        if not rows:
            return ExtractionResult(text="", chunks=[], sheet_count=1)

        headers = rows[0]
        header_str = " | ".join(headers)

        formatted_lines = [f"File: {filename}", f"Columns: {header_str}", f"Total Rows: {len(rows) - 1}\n"]

        # Batch rows into chunks of 35 rows to keep context tight and structured
        chunks: List[DocumentChunk] = []
        batch_size = 35

        for batch_start in range(1, len(rows), batch_size):
            batch_end = min(batch_start + batch_size, len(rows))
            batch_rows = rows[batch_start:batch_end]

            chunk_lines = [
                f"Columns: {header_str}",
                f"Rows {batch_start} to {batch_end - 1}:",
            ]
            for row in batch_rows:
                chunk_lines.append(" | ".join(row))

            chunk_text = "\n".join(chunk_lines)
            formatted_lines.extend(chunk_lines)

            chunks.extend(
                self.chunker.chunk_text(
                    chunk_text,
                    file_id=file_id,
                    filename=filename,
                    sheet="Data",
                    section=f"Rows {batch_start}–{batch_end - 1}",
                )
            )

        full_text = "\n".join(formatted_lines)
        return ExtractionResult(
            text=full_text,
            chunks=chunks,
            sheet_count=1,
            metadata={"headers": headers, "total_rows": len(rows) - 1},
        )

    def _extract_xlsx(self, file_id: str, filename: str, content_bytes: bytes) -> ExtractionResult:
        wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True, read_only=True)
        sheet_names = wb.sheetnames
        all_chunks: List[DocumentChunk] = []
        full_text_parts: List[str] = [f"Workbook: {filename}", f"Sheets: {', '.join(sheet_names)}\n"]

        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                # Clean row values
                row_vals = [str(v).strip() if v is not None else "" for v in row]
                if any(row_vals):  # Skip completely blank rows
                    rows.append(row_vals)

            if not rows:
                full_text_parts.append(f"Sheet: {sheet_name}\n(Empty Sheet)\n")
                continue

            headers = rows[0]
            header_str = " | ".join(headers)
            sheet_header = f"Sheet: {sheet_name}\nColumns: {header_str}\nTotal Rows: {len(rows) - 1}\n"
            full_text_parts.append(sheet_header)

            batch_size = 35
            for batch_start in range(1, len(rows), batch_size):
                batch_end = min(batch_start + batch_size, len(rows))
                batch_rows = rows[batch_start:batch_end]

                chunk_lines = [
                    f"Sheet: {sheet_name}",
                    f"Columns: {header_str}",
                    f"Rows {batch_start} to {batch_end - 1}:",
                ]
                for r in batch_rows:
                    chunk_lines.append(" | ".join(r))

                chunk_text = "\n".join(chunk_lines)
                all_chunks.extend(
                    self.chunker.chunk_text(
                        chunk_text,
                        file_id=file_id,
                        filename=filename,
                        sheet=sheet_name,
                        section=f"Sheet '{sheet_name}' Rows {batch_start}–{batch_end - 1}",
                    )
                )

        wb.close()
        return ExtractionResult(
            text="\n".join(full_text_parts),
            chunks=all_chunks,
            sheet_count=len(sheet_names),
            metadata={"sheets": sheet_names},
        )

    def _extract_xls(self, file_id: str, filename: str, content_bytes: bytes) -> ExtractionResult:
        try:
            wb = xlrd.open_workbook(file_contents=content_bytes)
            sheet_names = wb.sheet_names()
            all_chunks: List[DocumentChunk] = []
            full_text_parts: List[str] = [f"Workbook: {filename}", f"Sheets: {', '.join(sheet_names)}\n"]

            for sheet_name in sheet_names:
                ws = wb.sheet_by_name(sheet_name)
                if ws.nrows == 0:
                    continue

                rows = []
                for row_idx in range(ws.nrows):
                    row_vals = [str(ws.cell_value(row_idx, col_idx)).strip() for col_idx in range(ws.ncols)]
                    if any(row_vals):
                        rows.append(row_vals)

                if not rows:
                    continue

                headers = rows[0]
                header_str = " | ".join(headers)
                full_text_parts.append(f"Sheet: {sheet_name}\nColumns: {header_str}\nTotal Rows: {len(rows) - 1}\n")

                batch_size = 35
                for batch_start in range(1, len(rows), batch_size):
                    batch_end = min(batch_start + batch_size, len(rows))
                    batch_rows = rows[batch_start:batch_end]

                    chunk_lines = [
                        f"Sheet: {sheet_name}",
                        f"Columns: {header_str}",
                        f"Rows {batch_start} to {batch_end - 1}:",
                    ]
                    for r in batch_rows:
                        chunk_lines.append(" | ".join(r))

                    chunk_text = "\n".join(chunk_lines)
                    all_chunks.extend(
                        self.chunker.chunk_text(
                            chunk_text,
                            file_id=file_id,
                            filename=filename,
                            sheet=sheet_name,
                            section=f"Sheet '{sheet_name}' Rows {batch_start}–{batch_end - 1}",
                        )
                    )

            return ExtractionResult(
                text="\n".join(full_text_parts),
                chunks=all_chunks,
                sheet_count=len(sheet_names),
                metadata={"sheets": sheet_names},
            )
        except Exception as exc:
            logger.error(f"[XLS] Error parsing xls: {exc}")
            raise ValueError(f"Unable to parse legacy XLS spreadsheet: {exc}")

    def _extract_ods(self, file_id: str, filename: str, content_bytes: bytes) -> ExtractionResult:
        try:
            with zipfile.ZipFile(io.BytesIO(content_bytes)) as z:
                content_xml = z.read("content.xml")
                root = ET.fromstring(content_xml)
                all_chunks: List[DocumentChunk] = []
                full_text_parts: List[str] = [f"ODS Spreadsheet: {filename}\n"]

                # Namespace mapping
                ns = {"table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0", "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0"}
                tables = root.findall(".//table:table", ns)
                sheet_count = len(tables)

                for table in tables:
                    sheet_name = table.attrib.get("{urn:oasis:names:tc:opendocument:xmlns:table:1.0}name", "Sheet")
                    rows_data = []
                    for row in table.findall("table:table-row", ns):
                        cells = []
                        for cell in row.findall("table:table-cell", ns):
                            texts = [p.text for p in cell.findall("text:p", ns) if p.text]
                            cells.append(" ".join(texts) if texts else "")
                        if any(cells):
                            rows_data.append(cells)

                    if not rows_data:
                        continue

                    headers = rows_data[0]
                    header_str = " | ".join(headers)
                    full_text_parts.append(f"Sheet: {sheet_name}\nColumns: {header_str}\nTotal Rows: {len(rows_data) - 1}\n")

                    batch_size = 35
                    for batch_start in range(1, len(rows_data), batch_size):
                        batch_end = min(batch_start + batch_size, len(rows_data))
                        chunk_lines = [
                            f"Sheet: {sheet_name}",
                            f"Columns: {header_str}",
                            f"Rows {batch_start} to {batch_end - 1}:",
                        ]
                        for r in rows_data[batch_start:batch_end]:
                            chunk_lines.append(" | ".join(r))

                        chunk_text = "\n".join(chunk_lines)
                        all_chunks.extend(
                            self.chunker.chunk_text(
                                chunk_text,
                                file_id=file_id,
                                filename=filename,
                                sheet=sheet_name,
                                section=f"Sheet '{sheet_name}' Rows {batch_start}–{batch_end - 1}",
                            )
                        )

                return ExtractionResult(
                    text="\n".join(full_text_parts),
                    chunks=all_chunks,
                    sheet_count=sheet_count,
                )
        except Exception as exc:
            logger.error(f"[ODS] Error parsing ODS spreadsheet: {exc}")
            raise ValueError(f"Unable to parse OpenDocument Spreadsheet: {exc}")

