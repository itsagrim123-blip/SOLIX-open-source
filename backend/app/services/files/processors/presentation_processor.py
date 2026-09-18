import io
import logging
import os
import xml.etree.ElementTree as ET
import zipfile
from typing import List
from pptx import Presentation
from app.services.files.chunker import DocumentChunker
from app.services.files.models import DocumentChunk, ExtractionResult
from app.services.files.processors.base import BaseFileProcessor

logger = logging.getLogger("solix.files.presentation")


class PresentationProcessor(BaseFileProcessor):
    """Extracts slides, titles, bullet points, speaker notes, and tables from PPTX and ODP."""

    def __init__(self):
        self.chunker = DocumentChunker(chunk_size=1000, chunk_overlap=150)

    def can_process(self, category: str, mime_type: str, filename: str) -> bool:
        ext = os.path.splitext(filename.lower())[1]
        return category == "presentation" or ext in (".pptx", ".ppt", ".odp")

    async def extract(self, file_id: str, filename: str, content_bytes: bytes) -> ExtractionResult:
        ext = os.path.splitext(filename.lower())[1]

        if ext == ".ppt":
            logger.info(f"[PPT] Legacy binary .ppt detected for '{filename}'")
            raise ValueError(
                f"Solix can't read legacy binary '{filename}' format yet in this environment. "
                f"Please save or export the presentation as .pptx or .pdf."
            )

        if ext == ".odp":
            return self._extract_odp(file_id, filename, content_bytes)

        # PPTX extraction
        try:
            prs = Presentation(io.BytesIO(content_bytes))
        except Exception as exc:
            logger.error(f"[PPTX] Error opening presentation: {exc}")
            raise ValueError(f"Invalid or corrupted PPTX presentation: {exc}")

        slide_count = len(prs.slides)
        all_chunks: List[DocumentChunk] = []
        full_text_parts: List[str] = [f"Presentation: {filename}", f"Total Slides: {slide_count}\n"]

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_title = None

            # Check for slide title
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()

            slide_text_lines = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text and text != slide_title:
                            slide_text_lines.append(f"- {text}")
                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        if any(cells):
                            table_rows.append(" | ".join(cells))
                    if table_rows:
                        slide_text_lines.append("[Table]\n" + "\n".join(table_rows))

            # Extract speaker notes
            notes_text = None
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        notes_text = notes
            except Exception:
                pass

            # Build slide content string
            slide_header = f"Slide {slide_num}:"
            if slide_title:
                slide_header += f" {slide_title}"

            content_block = [slide_header]
            if slide_text_lines:
                content_block.append("\n".join(slide_text_lines))
            if notes_text:
                content_block.append(f"\nSpeaker Notes:\n{notes_text}")

            slide_full_str = "\n".join(content_block)
            full_text_parts.append(slide_full_str)

            chunks = self.chunker.chunk_text(
                slide_full_str,
                file_id=file_id,
                filename=filename,
                slide=slide_num,
                section=slide_title or f"Slide {slide_num}",
            )
            all_chunks.extend(chunks)

        return ExtractionResult(
            text="\n\n".join(full_text_parts),
            chunks=all_chunks,
            slide_count=slide_count,
            metadata={"total_slides": slide_count},
        )

    def _extract_odp(self, file_id: str, filename: str, content_bytes: bytes) -> ExtractionResult:
        try:
            with zipfile.ZipFile(io.BytesIO(content_bytes)) as z:
                content_xml = z.read("content.xml")
                root = ET.fromstring(content_xml)
                all_chunks: List[DocumentChunk] = []
                full_text_parts: List[str] = [f"ODP Presentation: {filename}\n"]

                ns = {
                    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
                    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
                    "presentation": "urn:oasis:names:tc:opendocument:xmlns:presentation:1.0",
                }
                pages = root.findall(".//draw:page", ns)
                slide_count = len(pages)

                for idx, page in enumerate(pages):
                    slide_num = idx + 1
                    texts = [p.text.strip() for p in page.findall(".//text:p", ns) if p.text and p.text.strip()]
                    slide_title = texts[0] if texts else f"Slide {slide_num}"
                    content_str = f"Slide {slide_num}: {slide_title}\n" + "\n".join(f"- {t}" for t in texts[1:])
                    full_text_parts.append(content_str)

                    chunks = self.chunker.chunk_text(
                        content_str,
                        file_id=file_id,
                        filename=filename,
                        slide=slide_num,
                        section=slide_title,
                    )
                    all_chunks.extend(chunks)

                return ExtractionResult(
                    text="\n\n".join(full_text_parts),
                    chunks=all_chunks,
                    slide_count=slide_count,
                )
        except Exception as exc:
            logger.error(f"[ODP] Error parsing ODP presentation: {exc}")
            raise ValueError(f"Unable to parse OpenDocument Presentation: {exc}")

